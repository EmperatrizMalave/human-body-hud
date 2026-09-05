from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor

# Crear documento
doc = docx.Document()

# Título Principal
title = doc.add_heading(
    "📘 MANUAL TÉCNICO DE USUARIO: MIGRACIÓN E INYECCIÓN DE DATOS", level=0
)
p_sub = doc.add_paragraph()
p_sub.add_run("Proyecto: ").bold = True
p_sub.add_run("Cyber-HUD Fitness App | ")
p_sub.add_run("Base de Datos: ").bold = True
p_sub.add_run("SQLite (gym.db) | ")
p_sub.add_run("Autor/a: ").bold = True
p_sub.add_run("Cori")

# Sección 1
doc.add_heading("1. INTRODUCCIÓN Y ARQUITECTURA DEL SISTEMA", level=1)
doc.add_paragraph(
    "Este manual describe el procedimiento estándar para convertir datasets estructurados en formato "
    "JSON e integrarlos dentro de una base de datos relacional SQLite (.db), utilizando un pipeline de migración "
    "en Python y administrando los datos mediante DBeaver."
)

# Sección 2
doc.add_heading(
    "2. PASO A PASO: PASAR DATA DE JSON A BASE DE DATOS .DB CON PYTHON", level=1
)

doc.add_heading("Paso 1: Creación del Script de Migración (cargar_db.py)", level=2)
code_block1 = """from pathlib import Path
import json
import sqlite3
import pandas as pd

BASE_DIR = Path(__file__).resolve().parent
json_path = BASE_DIR / "ejercicios.json"
db_path = BASE_DIR.parent / "backend" / "gym.db"

with open(json_path, "r", encoding="utf-8") as file:
    data = json.load(file)

df = pd.DataFrame(data)

conn = sqlite3.connect(db_path)
df.to_sql("ejercicios", conn, if_exists="replace", index=False)
conn.close()

print(f"¡Éxito! Carga completada en {db_path}. Se procesaron {len(df)} registros.")"""

p_code1 = doc.add_paragraph()
run1 = p_code1.add_run(code_block1)
run1.font.name = "Courier New"
run1.font.size = Pt(9.5)

doc.add_heading("Paso 2: Ejecución desde la Terminal", level=2)
doc.add_paragraph(
    "Para correr la migración desde la terminal con uv en Pop!_OS:"
)
p_term = doc.add_paragraph()
run_term = p_term.add_run("cd data\nuv run --with pandas cargar_db.py")
run_term.font.name = "Courier New"

doc.add_heading("Paso 3: Verificación en DBeaver", level=2)
doc.add_paragraph("1. Abrir DBeaver y conectar la base de datos ./backend/gym.db.")
doc.add_paragraph("2. Hacer clic derecho sobre la conexión y presionar Refrescar (F5).")
doc.add_paragraph("3. Abrir Tablas > ejercicios > Pestaña Data para validar los registros.")

# Sección 3
doc.add_heading("3. GUÍA PARA INYECCIÓN DE DATA", level=1)

doc.add_heading("Método A: Inyección Manual vía SQL (DBeaver)", level=2)
sql_code = """INSERT INTO ejercicios (
    TITLE_EN, TITLE_ES, MUSCLE_ID, RATINGDESC_ES, SENTADO_SITTING, LOW_IMPACT
) 
VALUES (
    'Cyber Plank Hold', 'Plancha isométrica Cyber', 'abs_core', 'Excelente', 0.0, 1.0
);"""
p_sql = doc.add_paragraph()
run_sql = p_sql.add_run(sql_code)
run_sql.font.name = "Courier New"

doc.add_heading("Método B: Inyección Masiva en Lote desde Python", level=2)
py_append = """import sqlite3
import pandas as pd
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
db_path = BASE_DIR.parent / "backend" / "gym.db"

nuevos_ejercicios = [
    {
        "TITLE_EN": "Neon Squat Jump",
        "TITLE_ES": "Sentadilla con salto Neón",
        "MUSCLE_ID": "quads_left",
        "RATINGDESC_ES": "Muy Bueno",
        "SENTADO_SITTING": 0.0,
        "LOW_IMPACT": 0.0
    }
]

df_nuevos = pd.DataFrame(nuevos_ejercicios)
conn = sqlite3.connect(db_path)
df_nuevos.to_sql("ejercicios", conn, if_exists="append", index=False)
conn.close()"""
p_app = doc.add_paragraph()
run_app = p_app.add_run(py_append)
run_app.font.name = "Courier New"

# Guardar documento
output_file = Path("../documentation/Manual_Tecnico_CyberHUD.docx")
output_file.parent.mkdir(parents=True, exist_ok=True)
doc.save(output_file)

print(f"¡Documento listo para Google Docs en: {output_file}!")